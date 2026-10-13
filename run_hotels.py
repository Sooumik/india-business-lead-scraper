"""Hotel & Restaurant lead scraper - Eastern India
(West Bengal, Odisha, Bihar, Jharkhand incl. Jamshedpur / Tatanagar; optional North-East).

Source : OpenStreetMap (Overpass API) - free and legal to use.
Enrich : visits each place's OWN website for a public email/phone (respects robots.txt;
         never visits Zomato / Booking.com / Facebook etc.).
Output : ONE tidy Excel workbook (+ a CSV) in the folder 'output_hotel_food'.

Needs only these files in the same folder:  run_hotels.py, enricher.py, config.py, requirements.txt

EXAMPLES
  python run_hotels.py                                    # hotels + restaurants, 4 eastern states
  python run_hotels.py --states Jharkhand --no-enrich     # quick test (Jamshedpur, Ranchi...)
  python run_hotels.py --sector hotel                     # hotels / guest houses only
  python run_hotels.py --sector restaurant                # restaurants / fast food / cafes only
  python run_hotels.py --group east,northeast             # add Sikkim, Assam, Tripura ...
  python run_hotels.py --no-enrich                        # faster, skips website visits
"""
import argparse
import math
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from tqdm import tqdm

import config
import enricher

# ------------------------------------------------------------------ SETTINGS
TARGET = 700
OUT_DIR = Path("output_hotel_food")
WORKERS = 8                    # parallel website visits (keep 8-10 or lower)

STATES = {
    "West Bengal":       {"iso": "IN-WB", "name": "West Bengal"},
    "Odisha":            {"iso": "IN-OD", "name": "Odisha|Orissa"},
    "Bihar":             {"iso": "IN-BR", "name": "Bihar"},
    "Jharkhand":         {"iso": "IN-JH", "name": "Jharkhand"},      # Jamshedpur = Tatanagar
    "Sikkim":            {"iso": "IN-SK", "name": "Sikkim"},
    "Assam":             {"iso": "IN-AS", "name": "Assam"},
    "Tripura":           {"iso": "IN-TR", "name": "Tripura"},
    "Meghalaya":         {"iso": "IN-ML", "name": "Meghalaya"},
    "Manipur":           {"iso": "IN-MN", "name": "Manipur"},
    "Nagaland":          {"iso": "IN-NL", "name": "Nagaland"},
    "Arunachal Pradesh": {"iso": "IN-AR", "name": "Arunachal Pradesh"},
    "Mizoram":           {"iso": "IN-MZ", "name": "Mizoram"},
}
GROUPS = {
    "east": ["West Bengal", "Odisha", "Bihar", "Jharkhand"],
    "northeast": ["Sikkim", "Assam", "Tripura", "Meghalaya", "Manipur", "Nagaland",
                  "Arunachal Pradesh", "Mizoram"],
    
}

# type key -> (sector, OSM selectors). One query per state + type keeps requests small.
TYPES = {
    "hotel":       ("Hotel",      ['["tourism"="hotel"]']),
    "guest_house": ("Hotel",      ['["tourism"="guest_house"]']),
    "other_stay":  ("Hotel",      ['["tourism"~"^(motel|hostel)$"]', '["leisure"="resort"]']),
    "restaurant":  ("Restaurant", ['["amenity"="restaurant"]']),
    "fast_food":   ("Restaurant", ['["amenity"="fast_food"]']),
    "cafe":        ("Restaurant", ['["amenity"="cafe"]']),
}
TYPE_LABEL = {"hotel": "Hotel", "guest_house": "Guest house / Lodge", "restaurant": "Restaurant",
              "fast_food": "Fast food", "cafe": "Cafe"}

OVERPASS_URLS = [
    "https://overpass-api.de/api/interpreter",
    "https://overpass.kumi.systems/api/interpreter",
]

# Never crawl these (aggregators / social sites): they are not the business's own site.
SKIP_DOMAINS = ("zomato", "swiggy", "booking.com", "tripadvisor", "makemytrip", "goibibo",
                "agoda", "expedia", "oyorooms", "oyo.", "facebook", "instagram", "twitter",
                "linkedin", "youtube", "google.", "justdial", "magicpin", "dineout",
                "wikipedia", "yelp", "airbnb", "trivago", "cleartrip", "ixigo", "yatra",
                "wa.me", "linktr.ee", "business.site")

# Major cities, used to fill a "Nearest city" column from the map position (within 25 km).
CITIES = {
    "Kolkata": (22.5726, 88.3639), "Howrah": (22.5958, 88.2636), "Siliguri": (26.7271, 88.3953),
    "Durgapur": (23.5204, 87.3119), "Asansol": (23.6739, 86.9524), "Darjeeling": (27.0360, 88.2627),
    "Kharagpur": (22.3460, 87.2320), "Haldia": (22.0667, 88.0698), "Bardhaman": (23.2324, 87.8615),
    "Malda": (25.0108, 88.1411), "Digha": (21.6275, 87.5093), "Kalimpong": (27.0594, 88.4695),
    "Jalpaiguri": (26.5167, 88.7333), "Cooch Behar": (26.3452, 89.4482), "Bankura": (23.2324, 87.0746),
    "Bhubaneswar": (20.2961, 85.8245), "Cuttack": (20.4625, 85.8830), "Puri": (19.8135, 85.8312),
    "Rourkela": (22.2604, 84.8536), "Sambalpur": (21.4669, 83.9812), "Berhampur": (19.3149, 84.7941),
    "Balasore": (21.4934, 86.9135), "Konark": (19.8876, 86.0945),
    "Patna": (25.5941, 85.1376), "Gaya": (24.7914, 85.0002), "Bhagalpur": (25.2425, 86.9842),
    "Muzaffarpur": (26.1209, 85.3647), "Darbhanga": (26.1542, 85.8918), "Purnia": (25.7771, 87.4753),
    "Arrah": (25.5560, 84.6603), "Bodh Gaya": (24.6961, 84.9869), "Hajipur": (25.6850, 85.2090),
    "Jamshedpur (Tatanagar)": (22.8046, 86.2029), "Ranchi": (23.3441, 85.3096),
    "Dhanbad": (23.7957, 86.4304), "Bokaro": (23.6693, 86.1511), "Hazaribagh": (23.9966, 85.3691),
    "Deoghar": (24.4764, 86.6954), "Gangtok": (27.3389, 88.6065), "Guwahati": (26.1445, 91.7362),
    "Agartala": (23.8315, 91.2868), "Shillong": (25.5788, 91.8933), "Imphal": (24.8170, 93.9368),
}

COLUMNS = ["name", "local_name", "sector", "type", "state", "city", "phone", "other_phones",
           "email", "other_emails", "website", "cuisine", "stars", "address", "lat", "lon"]


# ------------------------------------------------------------------ OVERPASS
def build_query(state_cfg, use_name_area, selectors):
    if use_name_area:
        head = ('area["boundary"="administrative"]["admin_level"="4"]'
                f'["name"~"^({state_cfg["name"]})$"]->.a;')
    else:
        head = f'area["ISO3166-2"="{state_cfg["iso"]}"]->.a;'
    body = "".join(f"nwr{s}(area.a);" for s in selectors)
    return "[out:json][timeout:300];" + head + "(" + body + ");out center tags;"


def overpass(query):
    """Returns list of elements, or None if every attempt failed."""
    for attempt in range(3):
        for url in OVERPASS_URLS:
            host = url.split("/")[2]
            try:
                r = requests.post(url, data={"data": query}, timeout=330,
                                  headers={"User-Agent": config.USER_AGENT})
                if r.status_code == 200:
                    return r.json().get("elements", [])
                print(f"   {host}: status {r.status_code} (busy), trying next...")
            except (requests.RequestException, ValueError) as e:
                print(f"   {host}: {type(e).__name__}, trying next...")
        wait = 10 * (attempt + 1)
        print(f"   waiting {wait}s before retry {attempt + 2}/3...")
        time.sleep(wait)
    return None


# ------------------------------------------------------------------ PARSING
def first(tags, *keys):
    for k in keys:
        if tags.get(k):
            return str(tags[k]).strip()
    return ""


def non_latin(s):
    return any(ord(c) > 0x024F for c in s)


def build_address(t):
    parts = [t.get("addr:housenumber"), t.get("addr:street"), t.get("addr:suburb"),
             t.get("addr:city"), t.get("addr:district"), t.get("addr:state"), t.get("addr:postcode")]
    return ", ".join(p for p in parts if p)


def parse(elements, state, type_key):
    sector = TYPES[type_key][0]
    rows = []
    for el in elements:
        t = el.get("tags", {})
        name = (t.get("name") or "").strip()
        if not name:
            continue
        local = ""
        if non_latin(name) and t.get("name:en"):
            local, name = name, t["name:en"].strip()
        if type_key == "other_stay":
            label = "Resort" if t.get("leisure") == "resort" else str(t.get("tourism", "stay")).title()
        else:
            label = TYPE_LABEL[type_key]
        rows.append({
            "name": name, "local_name": local, "sector": sector, "type": label, "state": state,
            "city": first(t, "addr:city", "addr:town", "addr:village"),
            "phone": first(t, "phone", "contact:phone", "mobile", "contact:mobile"),
            "other_phones": "",
            "email": first(t, "email", "contact:email").lower(),
            "other_emails": "",
            "website": first(t, "website", "contact:website", "url"),
            "cuisine": first(t, "cuisine").replace(";", ", ").replace("_", " "),
            "stars": first(t, "stars"),
            "address": build_address(t),
            "lat": el.get("lat") or el.get("center", {}).get("lat"),
            "lon": el.get("lon") or el.get("center", {}).get("lon"),
        })
    return rows


def last10(s):
    return re.sub(r"\D", "", str(s))[-10:]


def dedupe(rows):
    seen, out = set(), []
    for r in rows:
        nm = str(r["name"]).lower().strip()
        try:
            geo = (round(float(r["lat"]), 3), round(float(r["lon"]), 3))
        except (TypeError, ValueError):
            geo = None
        k1 = (nm, geo) if geo else None
        k2 = (nm, last10(r["phone"])) if last10(r["phone"]) else None
        if (k1 and k1 in seen) or (k2 and k2 in seen):
            continue
        if k1:
            seen.add(k1)
        if k2:
            seen.add(k2)
        out.append(r)
    return out


# ------------------------------------------------------------------ CLEANING
METRO2 = {"11", "20", "22", "33", "40", "44", "79", "80"}           # 2-digit area codes
STD3_7XX = {"712", "721", "731", "741", "751", "755", "761", "771"}  # 3-digit codes starting 7
METRO_STATE = {"80": "Karnataka", "79": "Gujarat"}
# Area codes (after the 0) that start with 6 in the eastern states
STD6 = {"Bihar": r"6(1\d|2\d|3\d|41|45|46|47)", "Jharkhand": r"6(43|5\d)", "Odisha": r"6[678]\d"}

EMAIL_RE = re.compile(r"^[a-z0-9._%+\-]+@[a-z0-9\-]+(\.[a-z0-9\-]+)*\.[a-z]{2,}$")
SCI_RE = re.compile(r"^-?\d(\.\d+)?E[+-]?\d+$", re.I)


def format_phone(part, state):
    """Returns (formatted, type) or None for junk."""
    p = str(part).strip()
    if not p or SCI_RE.match(p) or re.fullmatch(r"-\d+", p):
        return None
    if p.startswith("+") and not p.startswith("+91"):
        return (re.sub(r"\s+", " ", p), "International") if len(re.sub(r"\D", "", p)) >= 8 else None
    digits = re.sub(r"\D", "", p)
    if len(digits) < 6 or re.fullmatch(r"(\d)\1+", digits):
        return None
    if digits.startswith("1800") and len(digits) in (10, 11, 12):
        return (f"{digits[:4]} {digits[4:7]} {digits[7:]}", "Toll-free")
    if digits.startswith("91") and len(digits) in (12, 13):
        digits = digits[2:]
    leading_zero = digits.startswith("0")
    if leading_zero:
        digits = digits[1:]
    if len(digits) in (7, 8):
        return (f"(local) {digits}", "Landline (no area code)")
    if len(digits) != 10 or digits[0] == "0":
        return None

    if digits[0] in "12345":
        landline = True
    elif leading_zero:
        landline = (digits[:2] in {"79", "80"} or digits[:3] in STD3_7XX or
                    bool(state in STD6 and re.fullmatch(STD6[state], digits[:3])))
    elif digits[0] in "6789":
        landline = METRO_STATE.get(digits[:2]) == state
    else:
        landline = False

    if not landline:
        return (f"+91 {digits[:5]} {digits[5:]}", "Mobile")
    if digits[:2] in METRO2:
        return (f"+91 {digits[:2]} {digits[2:6]} {digits[6:]}", "Landline")
    return (f"+91 {digits[:3]} {digits[3:]}", "Landline")


def parse_phones(raw, state):
    out = []
    for part in re.split(r"\s*[;,]\s*|\s+/\s+", str(raw or "")):
        r = format_phone(part, state)
        if r and r[0] not in [x[0] for x in out]:
            out.append(r)
    return out


def clean_email(raw):
    e = str(raw or "").strip().lower().rstrip(".,;")
    return e if EMAIL_RE.match(e) else ""


def clean_website(raw):
    w = str(raw or "").strip().replace(" ", "").replace(",", ".")
    if not w or "@" in w:
        return ""
    if not w.startswith(("http://", "https://")):
        w = "https://" + w
    host = re.sub(r"^https?://(www\.)?", "", w).split("/")[0]
    return w if re.fullmatch(r"[a-z0-9\-]+(\.[a-z0-9\-]+)+", host, re.I) else ""


def nearest_city(lat, lon, limit_km=25):
    try:
        lat, lon = float(lat), float(lon)
    except (TypeError, ValueError):
        return ""
    best, best_d = "", limit_km
    for name, (a, b) in CITIES.items():
        dlat, dlon = math.radians(a - lat), math.radians(b - lon)
        h = math.sin(dlat / 2) ** 2 + math.cos(math.radians(lat)) * math.cos(math.radians(a)) * math.sin(dlon / 2) ** 2
        d = 6371 * 2 * math.asin(math.sqrt(h))
        if d < best_d:
            best, best_d = name, d
    return best


# ------------------------------------------------------------------ ENRICH
def skippable(url):
    host = urlparse(url if "//" in url else "//" + url).netloc.lower()
    return any(s in host for s in SKIP_DOMAINS)


def enrich_rows(rows, workers):
    targets = []
    for r in rows:
        site = clean_website(r["website"])
        if site and not skippable(site) and (not clean_email(r["email"]) or not r["phone"]):
            targets.append((r, site))
    print(f"\nVisiting {len(targets)} websites for emails/phones (polite crawl, can take a while)...")
    if not targets:
        return
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futures = {ex.submit(enricher.enrich_site, site): r for r, site in targets}
        for f in tqdm(as_completed(futures), total=len(futures)):
            r = futures[f]
            try:
                emails, phones = f.result()
            except Exception:
                continue
            if emails:
                if not clean_email(r["email"]):
                    r["email"] = emails[0]
                r["other_emails"] = ", ".join(e for e in emails if e != r["email"])
            if phones:
                known = last10(r["phone"])
                extra = [p for p in phones if last10(p) != known]
                if not r["phone"]:
                    r["phone"], extra = phones[0], phones[1:]
                r["other_phones"] = ", ".join(extra)


# ------------------------------------------------------------------ OUTPUT
OUT_COLS = ["No.", "Name", "Sector", "Type", "State", "Nearest City", "Area (as mapped)",
            "Phone 1", "Phone 1 Type", "Phone 2", "Other Phones", "Email", "Other Emails",
            "Website", "Cuisine", "Stars", "Address", "Local Name", "Map"]
WIDTHS = [6, 36, 12, 18, 14, 22, 18, 18, 18, 18, 24, 30, 26, 32, 20, 7, 45, 26, 40]


def finalize(rows):
    out = []
    for r in rows:
        state = r["state"]
        phones = parse_phones(r["phone"], state)
        for m in parse_phones(r["other_phones"], state):
            if m[0] not in [x[0] for x in phones]:
                phones.append(m)
        email = clean_email(r["email"])
        others = [e for e in (clean_email(x) for x in re.split(r"[;,\s]+", str(r["other_emails"]))) if e and e != email]
        lat, lon = r["lat"], r["lon"]
        has_geo = lat not in (None, "") and lon not in (None, "")
        out.append({
            "Name": r["name"], "Sector": r["sector"], "Type": r["type"], "State": state,
            "Nearest City": nearest_city(lat, lon), "Area (as mapped)": r["city"],
            "Phone 1": phones[0][0] if phones else "",
            "Phone 1 Type": phones[0][1] if phones else "",
            "Phone 2": phones[1][0] if len(phones) > 1 else "",
            "Other Phones": ", ".join(p[0] for p in phones[2:]),
            "Email": email, "Other Emails": ", ".join(others),
            "Website": clean_website(r["website"]),
            "Cuisine": r["cuisine"], "Stars": r["stars"], "Address": r["address"],
            "Local Name": r["local_name"],
            "Map": f"https://www.openstreetmap.org/?mlat={lat}&mlon={lon}#map=17/{lat}/{lon}" if has_geo else "",
        })
    df = pd.DataFrame(out)
    df = df.sort_values(["State", "Nearest City", "Sector", "Name"]).reset_index(drop=True)
    df.insert(0, "No.", range(1, len(df) + 1))
    return df[OUT_COLS]


def style(ws):
    fill = PatternFill("solid", start_color="1F4E78")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(WIDTHS[:ws.max_column], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.number_format = "@"


def write_outputs(final):
    has = (final["Phone 1"] != "") | (final["Email"] != "")
    with_contact = final[has].copy()
    with_contact["No."] = range(1, len(with_contact) + 1)

    summary = final.groupby(["State", "Sector"]).agg(
        Places=("Name", "count"),
        With_Phone=("Phone 1", lambda s: (s != "").sum()),
        With_Email=("Email", lambda s: (s != "").sum()),
        With_Website=("Website", lambda s: (s != "").sum()),
    ).reset_index()
    summary.columns = ["State", "Sector", "Places", "With Phone", "With Email", "With Website"]
    total = pd.DataFrame([["TOTAL", "", *summary.iloc[:, 2:].sum().tolist()]], columns=summary.columns)
    summary = pd.concat([summary, total], ignore_index=True)

    xlsx = OUT_DIR / "Hotel_Restaurant_Leads.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        final.to_excel(xw, sheet_name="All Leads", index=False)
        with_contact.to_excel(xw, sheet_name="With Contact", index=False)
        for sector, sheet in (("Hotel", "Hotels"), ("Restaurant", "Restaurants")):
            part = final[final.Sector == sector].copy()
            part["No."] = range(1, len(part) + 1)
            part.to_excel(xw, sheet_name=sheet, index=False)
        summary.to_excel(xw, sheet_name="Summary", index=False)
        for name in ("All Leads", "With Contact", "Hotels", "Restaurants"):
            style(xw.sheets[name])
        ws = xw.sheets["Summary"]
        for i, w in enumerate([20, 14, 10, 12, 12, 14], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c in ws[1]:
            c.font = Font(bold=True)
    final.to_csv(OUT_DIR / "Hotel_Restaurant_Leads.csv", index=False, encoding="utf-8-sig")
    return with_contact, summary


# ------------------------------------------------------------------ MAIN
def resolve(args):
    if args.states:
        states = [s.strip() for s in args.states.split(",") if s.strip()]
    else:
        states = []
        for g in [g.strip().lower() for g in args.group.split(",") if g.strip()]:
            if g == "all":
                g_list = list(GROUPS)
            elif g in GROUPS:
                g_list = [g]
            else:
                raise SystemExit(f"Unknown group '{g}'. Options: {list(GROUPS)} or 'all'")
            for gg in g_list:
                states += [s for s in GROUPS[gg] if s not in states]
    bad = [s for s in states if s not in STATES]
    if bad:
        raise SystemExit(f"Unknown state(s): {bad}. Options: {list(STATES)}")

    sectors = [s.strip().lower() for s in args.sector.split(",") if s.strip()]
    types = [k for k, (sec, _) in TYPES.items()
             if ("hotel" in sectors and sec == "Hotel") or ("restaurant" in sectors and sec == "Restaurant")]
    if not types:
        raise SystemExit("--sector must be hotel, restaurant or hotel,restaurant")
    return states, types


def main():
    ap = argparse.ArgumentParser(description="Hotel & Restaurant lead scraper (Eastern India)")
    ap.add_argument("--group", default="east", help="east, northeast, all (comma list). Default: east")
    ap.add_argument("--states", default="", help='exact names, e.g. --states "Odisha,Bihar" (overrides --group)')
    ap.add_argument("--sector", default="hotel,restaurant", help="hotel, restaurant, or both (default)")
    ap.add_argument("--no-enrich", action="store_true", help="skip website visits (faster)")
    args = ap.parse_args()

    states, types = resolve(args)
    print(f"States ({len(states)}): {', '.join(states)}")
    print(f"Types : {', '.join(types)}")

    OUT_DIR.mkdir(exist_ok=True)
    frames, failed = [], []

    for state in states:
        for tkey in types:
            raw = OUT_DIR / f"raw_{state.replace(' ', '_')}_{tkey}.csv"
            if raw.exists():
                print(f"[{state} / {tkey}] already saved, skipping")
                frames.append(pd.read_csv(raw, keep_default_na=False, dtype=str))
                continue
            print(f"\n=== {state} / {tkey} ===")
            cfg, sels = STATES[state], TYPES[tkey][1]
            elements = overpass(build_query(cfg, False, sels))
            if not elements:
                print("   nothing with state code, trying by state name...")
                elements = overpass(build_query(cfg, True, sels))
            if elements is None:
                print(f"!! {state} / {tkey}: server busy. Run the same command again in 10 minutes.")
                failed.append(f"{state}/{tkey}")
                continue
            rows = dedupe(parse(elements, state, tkey))
            print(f"   found {len(rows)}")
            part = pd.DataFrame(rows).reindex(columns=COLUMNS).fillna("")
            part.to_csv(raw, index=False, encoding="utf-8-sig")
            frames.append(part.astype(str))
            time.sleep(8)                      # be polite to the free server

    if not frames:
        raise SystemExit("\nNo data collected. Check internet / run again later.")

    rows = dedupe(pd.concat(frames, ignore_index=True).fillna("").to_dict("records"))
    if not args.no_enrich:
        enrich_rows(rows, WORKERS)

    final = finalize(rows)
    with_contact, summary = write_outputs(final)

    print("\n" + "=" * 56)
    print(f"TOTAL places       : {len(final)}   (target {TARGET})")
    print(f"With phone         : {(final['Phone 1'] != '').sum()}")
    print(f"With email         : {(final['Email'] != '').sum()}")
    print(f"With phone or email: {len(with_contact)}")
    print(summary.to_string(index=False))
    print(f"\nSaved in '{OUT_DIR}':  Hotel_Restaurant_Leads.xlsx  (open this one)  and .csv")
    if failed:
        print(f"\nNot finished: {', '.join(failed)} - run the same command again later.")
    if len(with_contact) < TARGET:
        print(f"\nOnly {len(with_contact)} have contact details. Add --group east,northeast for more states, "
              f"and don't use --no-enrich (website visits find extra emails/phones).")


if __name__ == "__main__":
    main()
